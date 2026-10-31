import sys
from pptx import Presentation
from lxml import etree
A='{http://schemas.openxmlformats.org/drawingml/2006/main}'
FILL={'0A0C0E':'FFFFFF','080D10':'FFFFFF','0B0D0F':'FFFFFF',
 '111518':'F4F5F3','121618':'F4F5F3','16191C':'F4F5F3',
 '2B312F':'E8EAE9','2D3435':'E8EAE9','37403F':'E3E6E5','3F4845':'E3E6E5',
 '7E898B':'BFC5C6','BBC2C4':'DDE1E2','FAFCFB':'FFFFFF'}
TEXT={'FAFCFB':'111111','BBC2C4':'2A3033','7E898B':'535C5E','879192':'535C5E','707C7D':'535C5E',
 'BEF02C':'4A7400','FF7D5E':'D2401F','37403F':'535C5E','3F4845':'535C5E'}
LINE={'BEF02C':'8DBA10'}
def kind(clr):
    par=clr.getparent().getparent()  # solidFill -> container
    t=etree.QName(par).localname
    if t in('rPr','defRPr','endParaRPr'): return 'text'
    if t=='ln': return 'line'
    return 'fill'
p=Presentation(sys.argv[1])
for s in p.slides:
    for clr in s._element.iter(A+'srgbClr'):
        if etree.QName(clr.getparent()).localname!='solidFill': continue
        v=clr.get('val').upper(); k=kind(clr)
        m={'text':TEXT,'line':{**FILL,**LINE},'fill':FILL}[k]
        if k=='text' and v in('0A0C0E','080D10','111518'): continue
        if v in m: clr.set('val',m[v])
p.save(sys.argv[2])
# 그림자 제거 (인쇄 시 회색 번짐)
p=Presentation(sys.argv[2])
P='{http://schemas.openxmlformats.org/presentationml/2006/main}'
for s in p.slides:
    for sp in s._element.iter(P+'sp'):
        spPr=sp.find(P+'spPr')
        if spPr is None or spPr.find(A+'effectLst') is not None: continue
        e=etree.SubElement(spPr,A+'effectLst')
        # 스키마 순서: ln 뒤
        ln=spPr.find(A+'ln')
        if ln is not None: ln.addnext(e)
p.save(sys.argv[2])
p=Presentation(sys.argv[2])
for s in p.slides:
    for er in s._element.iter(A+'effectRef'): er.set('idx','0')
p.save(sys.argv[2])
