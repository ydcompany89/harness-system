"""메가쇼 KDRS 집기 렌탈 요청서 (기본안 + 옵션, X배너 거치대 제외) → docx"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
HERE = os.path.dirname(os.path.abspath(__file__)); FONT = '맑은 고딕'
ITEMS = [  # 구역, 코드, 품명, 규격, 단가, 수량, 용도
 ('라라슈', 'KD-13', '타워쇼케이스', 'W900·D450·H2000', 140000, 1, '워킹화 3색 + 인솔 히어로 진열 (조명형)'),
 ('라라슈', 'KD-03', '3단진열대 2', 'W1480·D900·H1000', 110000, 1, '워킹화 사이즈·컬러 계단 진열'),
 ('라라슈', 'KD-02', '2단진열대 2', 'W1480·D600·H1000', 80000, 1, '인솔 M·L × 2색 체험대'),
 ('라라슈', 'KC-10', '벤치의자', 'W1450·D400·H400', 50000, 1, '시착 존'),
 ('라라슈', 'KC-03', '땅콩의자(검정)', 'W450·D580·H450', 20000, 2, '시착 보조 · 상담석'),
 ('라라슈', 'KET-13', '전신거울', '-', 30000, 1, '시착 존'),
 ('라라슈', 'KE-11', 'TV 50"', '스탠드 포함', 350000, 1, '브랜드 필름 · 원장 인터뷰 영상'),
 ('라라슈', 'KET-22', '아크릴단상 1', 'W600·H1160', 130000, 1, '정형외과 자문 코너 진열'),
 ('re.feeel', 'KD-13', '타워쇼케이스', 'W900·D450·H2000', 140000, 1, '양치컵 40° 연출 · CXP 원료 (조명형)'),
 ('re.feeel', 'KD-12', '3단쇼케이스진열대', 'W900·D450·H900', 70000, 1, '칫솔 · 칫솔형 치실 · 양치컵'),
 ('re.feeel', 'KD-01', '1단진열대 2', 'W1480·D600·H750', 60000, 1, '욕실화 매대'),
 ('re.feeel', 'KD-07', '선반진열대', 'W1000·D300', 30000, 3, '우측벽 욕실화 선반'),
 ('공용', 'KET-16', '스탠드카다로그꽂이', '-', 30000, 1, '다국어 리플렛'),
 ('공용', 'KET-14', 'POP거치대', 'W350·H1200~2000', 20000, 2, '가격 POP · 이벤트 POP'),
 ('공용', 'KET-09', '휴지통 2', '45L', 5000, 1, '카운터 뒤'),
]
sub = sum(i[4] * i[5] for i in ITEMS); vat = round(sub * 0.1); tot = sub + vat

d = Document(); sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7); sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = sec.bottom_margin = Cm(1.6)
st = d.styles['Normal']; st.font.name = FONT; st.font.size = Pt(9.5); st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
def shade(c, f):
    p = c._tc.get_or_add_tcPr(); s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), f); p.append(s)
def run(p, t, b=False, sz=9.5, col=None):
    r = p.add_run(t); r.bold = b; r.font.size = Pt(sz); r.font.name = FONT; r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    if col: r.font.color.rgb = RGBColor.from_string(col)
    return r
def para(t, sz=9.5, b=False, al=None, after=3, col=None):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    if al: p.alignment = al
    run(p, t, b, sz, col); return p
def cell(c, t, b=False, al=None, sz=9):
    c.text = ''; p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    if al: p.alignment = al
    run(p, t, b, sz)
def widths(t, cm):
    t.autofit = False; pr = t._tbl.tblPr; l = OxmlElement('w:tblLayout'); l.set(qn('w:type'), 'fixed'); pr.append(l)
    for i, g in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): g.set(qn('w:w'), str(int(cm[i] * 567)))
    for r in t.rows:
        for i, c in enumerate(r.cells): c.width = Cm(cm[i])

para('집기 렌탈 요청서', 18, True, WD_ALIGN_PARAGRAPH.CENTER, 2)
para('메가쇼 (킨텍스) 2026.11.12~15 · KDRS 카탈로그 기준', 10, False, WD_ALIGN_PARAGRAPH.CENTER, 10, '555555')
t = d.add_table(rows=4, cols=4); t.style = 'Table Grid'; widths(t, [3, 5.7, 3, 5.7])
info = [('수신', '㈜케이디알에스 (KDRS) · 02-907-9270', '발신', '㈜알앤디메이커스 신동규 대표'),
        ('전시회', '메가쇼 · 킨텍스', '전시 기간', '2026.11.12(목) ~ 11.15(일)'),
        ('부스', '기본부스 2개 (6m × 3m) · 부스번호 [  ]', '설치', '반입일 오전 (사무국 일정 확인 후 확정)'),
        ('연락처', '0507-1317-2516 · rndceo@rndmakers.kr', '작성일', '2026. 10.     .')]
for r, row in enumerate(info):
    for k, v in enumerate(row):
        cell(t.cell(r, k), v, b=(k % 2 == 0)); 
        if k % 2 == 0: shade(t.cell(r, k), 'EFEFEC')
para('', after=4)

para('1. 렌탈 품목', 12, True, after=3)
t = d.add_table(rows=1, cols=8); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
for k, h in enumerate(['구역', '코드', '품명', '규격', '단가', '수량', '금액', '용도']):
    cell(t.cell(0, k), h, True, WD_ALIGN_PARAGRAPH.CENTER); shade(t.cell(0, k), '0D0D0D')
    t.cell(0, k).paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string('C6F432')
for z, code, nm, spec, pr, q, use in ITEMS:
    c = t.add_row().cells
    vals = [z, code, nm, spec, f'{pr:,}', str(q), f'{pr*q:,}', use]
    for k, v in enumerate(vals):
        cell(c[k], v, al=WD_ALIGN_PARAGRAPH.RIGHT if k in (4, 6) else WD_ALIGN_PARAGRAPH.CENTER if k in (0, 1, 5) else None, sz=8)
    shade(c[0], {'라라슈': 'F1F7DC', 're.feeel': 'E3EDE6', '공용': 'F2F2F2'}[z])
for lab, v in [('공급가 합계', sub), ('부가세 (10%)', vat), ('합계 (부가세 포함)', tot)]:
    c = t.add_row().cells; m = c[0].merge(c[5]); cell(m, lab, True, WD_ALIGN_PARAGRAPH.RIGHT); cell(c[6], f'{v:,}', True, WD_ALIGN_PARAGRAPH.RIGHT, sz=8.5)
    for k in range(8): shade(c[k], 'F4F9E0' if lab.startswith('합계') else 'FAFAF7')
widths(t, [1.45, 1.45, 2.5, 3.1, 1.6, 0.9, 2.0, 4.4])
para(f'※ 표기 단가는 KDRS 카탈로그(부가세 별도) 기준 · 품목 15종 {sum(i[5] for i in ITEMS)}점', 8, col='666666', after=8)

para('2. 렌탈하지 않는 것 (참고)', 12, True, after=3)
for s in ['기본부스 제공: 안내데스크 1 · 의자 2 · 하이라이트 조명 → 렌탈 제외',
          'X배너 거치대: 출력물 업체(스터닝) 제작분 사용 → 렌탈 제외',
          '포디움 2: 자체 제작(허니콤 보드, 재사용) · 조명 보강: 사무국 스포트 추가 신청 + 클립 LED 자체 구매']:
    para('· ' + s, 9, after=1)
para('', after=4)

para('3. 확인 요청 사항', 12, True, after=3)
for i, s in enumerate(['위 단가가 전시 기간(4일) 전체 기준인지', '운반·설치·철거비 포함 여부 (별도라면 금액)',
                       'KD-07 선반진열대 설치 위치: 우측벽 판넬 3장에 1개씩 (높이 1,000 / 1,300 / 1,000mm 지그재그)',
                       'KD-13 타워쇼케이스 조명 포함 여부 · 전원(콘센트) 필요 수량',
                       'TV 50" 영상 재생 방식 (USB 재생 가능 여부)', '세금계산서 발행 (사업자: ㈜알앤디메이커스) · 견적서 회신 부탁드립니다'], 1):
    para(f'{i}) {s}', 9, after=1)
para('', after=4)
para('첨부: 부스 배치도 · 집기 위치도 (메가쇼_부스배치_집기_진열가이드.pdf)', 9, col='555555')
out = os.path.join(HERE, '메가쇼_집기렌탈_요청서_KDRS.docx'); d.save(out); print(out, sub, vat, tot)
