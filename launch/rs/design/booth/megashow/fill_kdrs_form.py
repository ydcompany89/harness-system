"""KDRS 비품임대신청서(원본 양식) 채우기 — 신청 10줄 → 품목 수만큼 확장. KDRS 카탈로그(10/16 수령본) 기준.
usage: python3 -I fill_kdrs_form.py <원본양식.xlsx> <출력.xlsx>"""
import sys, copy
import openpyxl
from openpyxl.worksheet.cell_range import CellRange
src, out = sys.argv[1], sys.argv[2]
ITEMS = [  # 제품명, 규격, 수량, 단가, 비고
 ('타워쇼케이스진열대(냉장냉동X)', 'W900×450×2000', 2, 140000, '라라슈 1 · re.feeel 1'),
 ('3단진열대', 'W1480×900×1000', 1, 110000, '워킹화'),
 ('2단진열대', 'W1480×600×1000', 1, 80000, '인솔 체험'),
 ('땅콩의자(검정)', 'W450×580×450', 4, 20000, '시착 3 · 상담 1'),
 ('전신거울', '-', 1, 30000, '시착'),
 ('TV 50"', '-', 1, 350000, '영상 USB 재생'),
 ('1단쇼케이스진열대(냉장냉동X)', 'W900×450×900', 1, 60000, '자문 코너'),
 ('3단쇼케이스진열대(냉장냉동X)', 'W900×450×900', 1, 70000, '구강용품'),
 ('1단진열대', 'W1480×600×750', 1, 60000, '욕실화'),
 ('선반진열대(옥타부스용)', 'W1000×300', 3, 30000, '우측벽 판넬 각 1'),
 ('스탠드 카달로그 꽂이', '-', 1, 30000, '리플렛'),
 ('POP스텐드', 'W350×1200~2000', 2, 20000, 'POP'),
 ('행운함', 'W350×350×350', 1, 25000, '명함 응모함'),
 ('휴지통', '50L', 1, 5000, ''),
]
wb = openpyxl.load_workbook(src)
ws = wb['비품임대신청서']
N = len(ITEMS); EXTRA = max(0, N - 10)
if EXTRA:
    merges = [CellRange(str(m)) for m in ws.merged_cells.ranges]
    for m in merges:
        if m.max_row >= 19: ws.unmerge_cells(str(m))
    heights = {r: ws.row_dimensions[r].height for r in range(19, 40)}
    ws.insert_rows(19, EXTRA)
    for r, h in heights.items(): ws.row_dimensions[r + EXTRA].height = h
    for i in range(EXTRA):
        r = 19 + i; ws.row_dimensions[r].height = ws.row_dimensions[18].height
        for col in 'ABCDEFGH':
            ws[f'{col}{r}']._style = copy.copy(ws[f'{col}18']._style)
        ws.merge_cells(f'B{r}:C{r}')
    for m in merges:
        if m.max_row < 19: continue
        if m.min_row >= 19: m.shift(row_shift=EXTRA)
        else: m.expand(down=EXTRA)
        ws.merge_cells(str(m))
    ws.print_area = f'A1:Q{36 + EXTRA}'
first, last = 9, 8 + N
for i, (nm, spec, q, p, note) in enumerate(ITEMS):
    r = first + i
    ws[f'A{r}'] = i + 1; ws[f'B{r}'] = nm; ws[f'D{r}'] = spec; ws[f'E{r}'] = q; ws[f'F{r}'] = p
    ws[f'G{r}'] = f'=E{r}*F{r}'; ws[f'H{r}'] = note
    ws[f'G{r}'].number_format = '#,##0"원"'; ws[f'F{r}'].number_format = '#,##0'
    for col in 'BDH':
        a = copy.copy(ws[f'{col}{r}'].alignment); a.shrink_to_fit = True; a.wrap_text = False; ws[f'{col}{r}'].alignment = a
tot = last + 1
ws[f'F{tot}'] = f'=SUM(G{first}:G{last})'; ws[f'F{tot+1}'] = f'=ROUND(F{tot}*0.1,0)'; ws[f'F{tot+2}'] = f'=F{tot}+F{tot+1}'
for r in (tot, tot + 1, tot + 2): ws[f'F{r}'].number_format = '#,##0"원"'
ws['C2'] = '메가쇼 (킨텍스) 2026.11.12~15'; ws['G2'] = ''
ws['C3'] = '신동규'; ws['G3'] = '2026.10.   .'
ws['C4'] = '㈜알앤디메이커스            (인)'; ws['G4'] = '신동규'
ws['C5'] = '0507-1317-2516'
ws['C6'] = '010-6880-2516'; ws['G6'] = 'rndceo@rndmakers.kr'
ws[f'A{27 + EXTRA}'] = '■세금계산서 발행             □거래명세표 발행           □입금표 발행           □기타'
wb.save(out); print('items', N, 'rows', first, last, 'total row', tot)
