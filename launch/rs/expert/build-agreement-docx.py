"""doctor-collab-plan.md §6 계약서 → doctor-agreement-draft.docx (python3 build-agreement-docx.py)"""
import re, os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
md = open(os.path.join(HERE, 'doctor-collab-plan.md'), encoding='utf-8').read()
body = md.split('# 제품 설계 자문 및 초상·성명 이용 계약서', 1)[1]

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.2); s.left_margin = s.right_margin = Cm(2.5)
st = doc.styles['Normal']; st.font.name = 'Malgun Gothic'; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.3

def runs(p, text):
    for i, part in enumerate(re.split(r'\*\*', text)):
        if part: r = p.add_run(part); r.bold = (i % 2 == 1)

t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('제품 설계 자문 및 초상·성명 이용 계약서'); r.bold = True; r.font.size = Pt(17)
doc.add_paragraph()
table_rows = []
for line in body.strip().splitlines():
    line = line.rstrip()
    if not line: continue
    if line.startswith('|'):
        cells = [c.strip() for c in line.strip('|').split('|')]
        if not set(''.join(cells)) <= set('-: '): table_rows.append(cells)
        continue
    if table_rows:
        tb = doc.add_table(rows=len(table_rows), cols=2); tb.style = 'Table Grid'
        for i, row in enumerate(table_rows):
            for j, c in enumerate(row): runs(tb.cell(i, j).paragraphs[0], c)
        table_rows = []; doc.add_paragraph()
    if line.startswith('>'):
        p = doc.add_paragraph(); runs(p, line.lstrip('> ')); p.runs[0].italic = True; continue
    p = doc.add_paragraph()
    if re.match(r'^\d+\. ', line): p.paragraph_format.left_indent = Cm(0.8)
    if line.startswith('2026년'): p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if line.startswith('**제') : p.paragraph_format.space_before = Pt(8)
    runs(p, line)
if table_rows:
    tb = doc.add_table(rows=len(table_rows), cols=2); tb.style = 'Table Grid'
    for i, row in enumerate(table_rows):
        for j, c in enumerate(row): runs(tb.cell(i, j).paragraphs[0], c)
out = os.path.join(HERE, 'doctor-agreement-draft.docx'); doc.save(out); print('saved', out)
