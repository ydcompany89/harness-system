"""2차 사업계획 변경 신청서 (별지 제1-3호 서식) docx 생성 — 1차(260917) 양식 구조 그대로."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

FONT = "맑은 고딕"
OUT = os.path.join(os.path.dirname(__file__), "change2-request-2026-10.docx")

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2)
sec.top_margin = sec.bottom_margin = Cm(1.8)
st = d.styles["Normal"]
st.font.name = FONT; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

def shade(cell, hexcolor="E7E6E6"):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcolor)
    tcPr.append(s)

def put(cell, lines, bold=False, center=False, size=10, color=None):
    cell.text = ""
    if isinstance(lines, str): lines = [lines]
    for i, ln in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        b = bold
        if ln.startswith("**"): ln = ln.strip("*"); b = True
        r = p.add_run(ln); r.bold = b; r.font.size = Pt(size); r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        if color: r.font.color.rgb = RGBColor.from_string(color)

def widths(t, cm):
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for i, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
        if i < len(cm): gc.set(qn("w:w"), str(int(cm[i] * 567)))
    for row in t.rows:
        for i, c in enumerate(row.cells):
            if i < len(cm): c.width = Cm(cm[i])

def para(text, size=10, bold=False, align=None, after=4):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    if align: p.alignment = align
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return p

# ── 1쪽: 변경 신청서 ──
para("[별지 제1-3호 서식]", 9)
para("새활용 산업 육성 지원사업", 13, True, WD_ALIGN_PARAGRAPH.CENTER, 0)
para("사업계획 변경 신청서 (2차)", 18, True, WD_ALIGN_PARAGRAPH.CENTER, 10)

t = d.add_table(rows=2, cols=4); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
widths(t, [3.0, 5.5, 3.0, 5.5])
put(t.cell(0,0), "지원기업명", True, True); shade(t.cell(0,0))
put(t.cell(0,1), "알앤디메이커스(주)")
put(t.cell(0,2), "사업 책임자", True, True); shade(t.cell(0,2))
put(t.cell(0,3), "신동규 대표")
put(t.cell(1,0), "1차 변경 승인", True, True); shade(t.cell(1,0))
c = t.cell(1,1).merge(t.cell(1,3)); put(c, "2026. 9. 18. (환산협2026-1079) — 본 신청서의 ‘변경 전’은 1차 변경 승인 내용 기준")
para("", after=2)

BEFORE = [
    "**8. 각종 수수료 및 사용료 (총 12,173천원)**",
    "- KCL 시험 수수료 4,000천원",
    "- 환경성적표지·저탄소 인증 수수료 2,000천원",
    "- KC 인증 시험 수수료 2,000천원",
    "- 특허·디자인권 출원/등록 수수료 4,173천원",
]
AFTER = [
    "**8. 각종 수수료 및 사용료 (총 12,173천원)**",
    "- KCL 미끄럼 방지 및 안전성 시험 수수료 4,000천원",
    "- 환경성적표지·저탄소 인증 수수료 2,000천원",
    "- KC 인증 시험 수수료 2,000천원",
    "- 지식재산권(특허·디자인·상표) 출원/등록 수수료 4,173천원",
    "  · 재생소재 상표 「리타이어렌(RETIRENE)」 출원(폐타이어 재생고무·폐PVC 배합 소재, 국문·영문 / 제17·25류) 1,248천원",
    "  · 특허 「장시간 기립 작업자의 하지 및 허리 피로 저감을 위한 보행 궤적 배치형 미끄럼 방지 신발」 출원·우선심사 2,600천원",
    "  · 출원 전 선행기술·유사상표 조사 325천원",
]
REASON = [
    "각종 수수료 및 사용료: 폐타이어 재생고무와 폐PVC(실크벽지)를 배합한 재생소재의 브랜드화(상표 「리타이어렌」) 및 장시간 기립 작업자용 미끄럼 방지 신발 특허 출원 계획이 구체화됨에 따라, 지식재산권 내역의 범위를 ‘특허·디자인권’에서 ‘특허·디자인·상표’로 확대하고 세부 내역을 명시함",
    "※ 세목·사업용도 금액 변동 없음 (각종 수수료 및 사용료 12,173천원 유지, 사업비 총계 177,200천원 유지)",
]

t = d.add_table(rows=3, cols=3); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(["구분", "변경 전", "변경 후"]):
    put(t.cell(0,i), h, True, True); shade(t.cell(0,i))
put(t.cell(1,0), "변경내용", True, True); shade(t.cell(1,0))
put(t.cell(1,1), BEFORE, size=9)
put(t.cell(1,2), AFTER, size=9)
put(t.cell(2,0), "변경사유", True, True); shade(t.cell(2,0))
c = t.cell(2,1).merge(t.cell(2,2)); put(c, REASON, size=9)
widths(t, [2.4, 6.6, 8.0])

para("", after=2)
para("위와 같이 사업계획 변경을 신청합니다.", 10, after=2)
para("붙임  1. 사업비 소요명세서 변경사항 1부.", 10, after=0)
para("      2. 변경 사업계획서(수정본) 1부.  끝.", 10, after=10)
para("2026.   10.     .", 11, False, WD_ALIGN_PARAGRAPH.CENTER, 6)
para("신청인 : 알앤디메이커스(주) 대표 신 동 규   (인)", 11, False, WD_ALIGN_PARAGRAPH.RIGHT, 14)
para("운영기관장 귀하", 13, True, WD_ALIGN_PARAGRAPH.LEFT, 0)

# ── 2쪽: 사업비 소요명세서 변경사항 ──
d.add_page_break()
para("사업비 소요명세서 변경사항", 15, True, WD_ALIGN_PARAGRAPH.CENTER, 2)
para("(단위 : 천원)", 9, False, WD_ALIGN_PARAGRAPH.RIGHT, 2)
ROWS = [
    # 보조비목, 보조세목, 사업용도, 전, 후, 증감
    ("사업비 총계", "", "", "177,200", "177,200", "-"),
    ("인건비", "보수", "참여자 인건비", "26,580", "26,580", "-"),
    ("", "일용임금", "단기채용 인건비", "X", "X", "-"),
    ("운영비", "재료비", "기자재 구입", "-", "-", "-"),
    ("", "", "부품·재료 구입", "13,000", "13,000", "-"),
    ("", "일반용역비", "시제품 외주제작", "22,328", "22,328", "-"),
    ("", "", "설비 외주제작", "49,100", "49,100", "-"),
    ("", "", "홍보 외주제작", "10,700", "10,700", "-"),
    ("", "", "전문가 활용비", "4,700", "4,700", "-"),
    ("", "일반수용비", "안내홍보물 제작비", "11,000", "11,000", "-"),
    ("", "", "소모성 물품 구입비", "129", "129", "-"),
    ("", "", "각종 수수료 및 사용료 ※내역 변경", "12,173", "12,173", "- (내역 변경)"),
    ("", "", "행사 지원 경비", "9,770", "9,770", "-"),
    ("", "", "홍보비", "17,720", "17,720", "-"),
    ("여비", "국내여비", "국내여비", "-", "-", "-"),
    ("", "국외여비", "국외여비", "0", "0", "-"),
]
t = d.add_table(rows=len(ROWS)+1, cols=6); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
widths(t, [2.2, 2.4, 4.8, 2.6, 2.6, 2.4])
for i, h in enumerate(["보조비목", "보조세목", "사업용도", "변경 전 사업비\n(1차 승인)", "변경 후 사업비", "증감"]):
    put(t.cell(0,i), h.split("\n"), True, True, 9); shade(t.cell(0,i))
for r, row in enumerate(ROWS, 1):
    hl = "수수료" in row[2]
    for c, v in enumerate(row):
        put(t.cell(r,c), v, bold=(r == 1 or hl), center=True, size=9)
        if hl: shade(t.cell(r,c), "FFF2CC")
    if r == 1: [shade(t.cell(1,c), "F2F2F2") for c in range(6)]
para("※ 사업비 내역은 현물, 현금 관계없이 변경 요청 이전 금액과 변경희망 금액을 작성합니다", 8, after=6)

para("각종 수수료 및 사용료 — 내역 대비표", 11, True, after=2)
CMP = [
    ("KCL 시험 수수료", "4,000", "KCL 미끄럼 방지 및 안전성 시험 수수료", "4,000"),
    ("환경성적표지·저탄소 인증 수수료", "2,000", "환경성적표지·저탄소 인증 수수료", "2,000"),
    ("KC 인증 시험 수수료", "2,000", "KC 인증 시험 수수료", "2,000"),
    ("특허·디자인권 출원/등록 수수료", "4,173", "지식재산권(특허·디자인·상표) 출원/등록 수수료", "4,173"),
    ("", "", "  ㄴ 재생소재 상표 「리타이어렌(RETIRENE)」 출원 (국·영문, 17·25류)", "1,248"),
    ("", "", "  ㄴ 보행 궤적 배치형 미끄럼 방지 신발 특허 출원·우선심사", "2,600"),
    ("", "", "  ㄴ 선행기술·유사상표 조사", "325"),
    ("합계", "12,173", "합계", "12,173"),
]
t = d.add_table(rows=len(CMP)+1, cols=4); t.style = "Table Grid"
for i, h in enumerate(["변경 전 내역", "금액", "변경 후 내역", "금액"]):
    put(t.cell(0,i), h, True, True, 9); shade(t.cell(0,i))
for r, row in enumerate(CMP, 1):
    for c, v in enumerate(row):
        put(t.cell(r,c), v, bold=(row[0] == "합계"), center=(c in (1,3)), size=9)
widths(t, [5.0, 1.8, 8.4, 1.8])

d.save(OUT); print(OUT)
